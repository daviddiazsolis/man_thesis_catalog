"""Exporta el catálogo (src/data/theses/*.json + advisors.json) a un Excel de respaldo.

Uso (desde la raíz del repo):  python gen/exportar_excel.py [ruta_salida.xlsx]
Por defecto escribe  catalogo_tesis_MAN.xlsx  en la raíz del repo.
Requiere: pip install openpyxl
"""
import json, glob, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(ROOT, 'catalogo_tesis_MAN.xlsx')
adv = json.load(open(os.path.join(ROOT, 'src/data/advisors.json'), encoding='utf-8'))
theses = [json.load(open(f, encoding='utf-8')) for f in sorted(glob.glob(os.path.join(ROOT, 'src/data/theses/*.json')))]
theses.sort(key=lambda t: (t['defenseDate'], t['student']), reverse=True)
name = lambda k: adv[k]['name'] if k in adv else k
j = lambda xs: '; '.join(xs)

HEAD = Font(bold=True, color='FFFFFF'); FILL = PatternFill('solid', fgColor='0034D9')

def sheet(ws, header, rows, widths):
    ws.append(header)
    for c in ws[1]: c.font = HEAD; c.fill = FILL; c.alignment = Alignment(vertical='center', wrap_text=True)
    for r in rows: ws.append(r)
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(vertical='top', wrap_text=True)
    ws.freeze_panes = 'C2'

wb = Workbook()
ws = wb.active; ws.title = 'Catalogo'
sheet(ws,
  ['ID', 'Estado', 'Autorización', 'Alumno/a', 'Generación', 'Programa', 'Fecha defensa', 'Año defensa',
   'Profesor guía', 'Profesor co-guía', 'Comisión', 'Título (ES)', 'Título (EN)', 'Sector', 'Tipo de problema',
   'Técnicas', 'Herramientas', 'Datasets', 'Datos públicos', 'Organización', 'Impacto social (nivel)',
   'Impacto social (ES)', 'Impacto social (EN)', 'Palabras clave (ES)', 'Palabras clave (EN)',
   'Resumen ejecutivo (ES)', 'Executive summary (EN)',
   'Highlight 1 (ES)', 'Highlight 2 (ES)', 'Highlight 3 (ES)', 'Highlight 4 (ES)', 'Highlight 5 (ES)',
   'Highlight 1 (EN)', 'Highlight 2 (EN)', 'Highlight 3 (EN)', 'Highlight 4 (EN)', 'Highlight 5 (EN)',
   'URL repositorio U. de Chile', 'PDF en el sitio', 'Notas'],
  [[t['id'], t['status'], t['authorization'], t['student'], t['cohort'], t['program'], t['defenseDate'], t['defenseYear'],
    name(t['advisor']), name(t['coAdvisor']) if t['coAdvisor'] else '', j([name(c) for c in t['committee']]),
    t['title']['es'], t['title']['en'], j(t['domain']), j(t['taskTypes']), j(t['techniques']), j(t['tools']),
    j([f"{d['nameEs']} ({d['source']})" if d.get('source') else d['nameEs'] for d in t['datasets']]),
    'Sí' if any(d['public'] for d in t['datasets']) else 'No', t['organization'] or '',
    t['socialImpact']['level'], t['socialImpact']['es'], t['socialImpact']['en'],
    j(t['keywords']['es']), j(t['keywords']['en']), t['abstract']['es'], t['abstract']['en'],
    *(t['highlights']['es'] + [''] * 5)[:5], *(t['highlights']['en'] + [''] * 5)[:5],
    t['repositoryUrl'] or '', ('https://man-thesis-catalog.vercel.app' + t['pdfUrl']) if t.get('pdfUrl') else '', t.get('notes') or ''] for t in theses],
  [16, 9, 12, 28, 10, 9, 12, 8, 26, 22, 30, 50, 50, 20, 26, 40, 24, 50, 10, 30, 12, 45, 45, 35, 35, 80, 80] + [45] * 10 + [40, 40, 40])

ws2 = wb.create_sheet('Datasets')
sheet(ws2, ['ID tesis', 'Alumno/a', 'Dataset (ES)', 'Dataset (EN)', 'Fuente', 'Público', 'URL', 'Tamaño'],
  [[t['id'], t['student'], d['nameEs'], d['nameEn'], d.get('source') or '', 'Sí' if d['public'] else 'No', d.get('url') or '', d.get('sizeNote') or '']
   for t in theses for d in t['datasets']], [16, 28, 45, 45, 40, 8, 40, 30])

ws3 = wb.create_sheet('Profesores')
sheet(ws3, ['Clave', 'Nombre', 'Afiliación (ES)', 'Afiliación (EN)', 'URL', 'Tipo de URL', 'Tesis como guía', 'Tesis en comisión'],
  [[k, a['name'], a['affiliationEs'], a['affiliationEn'], a.get('url') or '', a.get('urlType') or '',
    sum(1 for t in theses if t['advisor'] == k or t['coAdvisor'] == k), sum(1 for t in theses if k in t['committee'])]
   for k, a in adv.items()], [14, 36, 60, 60, 60, 12, 12, 12])

ws4 = wb.create_sheet('Tecnicas')
from collections import Counter
cnt = Counter(x for t in theses for x in t['techniques'])
sheet(ws4, ['Técnica', 'N tesis', 'Tesis'], [[k, v, j([t['id'] for t in theses if k in t['techniques']])] for k, v in cnt.most_common()], [40, 8, 80])

wb.save(OUT)
print(f'Excel escrito: {OUT}  ({len(theses)} tesis)')
